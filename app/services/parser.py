"""Multi-format document parser.

Supported formats: PDF, TXT, MD, DOCX, PPTX, XLSX.
Each parser returns: list[dict] with keys:
    - text: str
    - source: str (filename)
    - page: int or str (page number, slide number, or sheet name)
"""

import os
from pathlib import Path
from typing import List, Dict, Any

try:
    import pymupdf as fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def parse_document(filepath: str) -> List[Dict[str, Any]]:
    """Route document to appropriate parser based on file extension.

    Args:
        filepath: Absolute or relative path to the document file.

    Returns:
        List of dicts with 'text', 'source', 'page' keys.
    """
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {filepath}")

    ext = path.suffix.lower()
    parsers = {
        ".pdf": _parse_pdf,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".docx": _parse_docx,
        ".pptx": _parse_pptx,
        ".xlsx": _parse_xlsx,
    }

    parser = parsers.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file format: '{ext}'. Supported: {list(parsers.keys())}")

    return parser(filepath)


def _parse_pdf(filepath: str) -> List[Dict[str, Any]]:
    """Extract text from PDF page by page."""
    if fitz is None:
        raise ImportError("PyMuPDF is not installed.")

    pages = []
    filename = Path(filepath).name
    with fitz.open(filepath) as doc:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text()
            if text and text.strip():
                pages.append({
                    "text": text.strip(),
                    "source": filename,
                    "page": page_num,
                })

    if not pages:
        raise ValueError(f"No extractable text found in '{filename}'. It may be scanned or empty.")
    return pages


def _parse_text(filepath: str) -> List[Dict[str, Any]]:
    """Extract text from plain text or markdown file."""
    filename = Path(filepath).name
    # Attempt reading with utf-8, fallback to latin-1
    try:
        text = Path(filepath).read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = Path(filepath).read_text(encoding="latin-1", errors="replace")

    if not text.strip():
        raise ValueError(f"File '{filename}' is empty.")

    return [{
        "text": text.strip(),
        "source": filename,
        "page": 1,
    }]


def _parse_docx(filepath: str) -> List[Dict[str, Any]]:
    """Extract text from DOCX file."""
    if DocxDocument is None:
        raise ImportError("python-docx is not installed.")

    filename = Path(filepath).name
    doc = DocxDocument(filepath)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract tables if present
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)
    if not full_text.strip():
        raise ValueError(f"No text found in Word document '{filename}'.")

    return [{
        "text": full_text.strip(),
        "source": filename,
        "page": 1,
    }]


def _parse_pptx(filepath: str) -> List[Dict[str, Any]]:
    """Extract text from PowerPoint slides."""
    if Presentation is None:
        raise ImportError("python-pptx is not installed.")

    filename = Path(filepath).name
    prs = Presentation(filepath)
    slides = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    p_text = "".join(run.text for run in paragraph.runs).strip()
                    if p_text:
                        slide_texts.append(p_text)
        if slide_texts:
            slides.append({
                "text": "\n".join(slide_texts),
                "source": filename,
                "page": slide_idx,
            })

    if not slides:
        raise ValueError(f"No text found in presentation '{filename}'.")
    return slides


def _parse_xlsx(filepath: str) -> List[Dict[str, Any]]:
    """Extract tabular text from Excel spreadsheet sheets."""
    if load_workbook is None:
        raise ImportError("openpyxl is not installed.")

    filename = Path(filepath).name
    wb = load_workbook(filepath, read_only=True, data_only=True)
    sheets = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        row_lines = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip() != ""]
            if cells:
                row_lines.append(" | ".join(cells))

        if row_lines:
            sheets.append({
                "text": f"Sheet: {sheet_name}\n" + "\n".join(row_lines),
                "source": filename,
                "page": sheet_name,
            })

    wb.close()
    if not sheets:
        raise ValueError(f"No text or data found in spreadsheet '{filename}'.")
    return sheets
